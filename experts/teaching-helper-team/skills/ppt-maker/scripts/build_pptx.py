#!/usr/bin/env python3
"""build_pptx.py - traditional .pptx generator (python-pptx) for the ppt-maker skill.

Fallback path used only when the deliverable MUST be a native PowerPoint file
(LMS requires .pptx / PowerPoint-only classroom / handoff to PPT-only colleagues).
Reads the SAME deck spec JSON as build_deck.py. Requires the optional dependency
python-pptx (optional; see requirements.lock at the package root; installed into the team's managed venv by setup).
Run via $VENV_PY (never bare python3).
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
except ImportError:
    # venv 校验：未用托管 venv 运行时提示更明确的解释器路径
    venv_hint = ""
    if "teaching-helper-team" not in str(sys.executable):
        venv_hint = (f"\n  当前解释器: {sys.executable}\n"
                     "  提示: 请用 setup-status.json 中的 venv_python 运行（禁止裸 python3），"
                     "如: $VENV_PY build_pptx.py --spec deck.json --out deck.pptx")
    sys.stderr.write(
        "ERROR: python-pptx 未安装。传统 .pptx 路径为可选依赖；\n"
        "请先运行 skills/setup 安装，或改用 HTML 方案（build_deck.py，零依赖）。"
        + venv_hint + "\n"
    )
    raise SystemExit(3)  # 统一错误码：3 = 缺依赖（与 extract_pptx/extract_pdf 一致）

# Preset palette (hex) — mirrors build_deck.py PRESETS for visual continuity.
PRESET_COLORS = {
    "academic-light": ("#1f3a5f", "#e8833a", "#ffffff", "#1f2733"),
    "ink": ("#23201c", "#9e2b25", "#f7f3ea", "#23201c"),
    "aurora": ("#37e0c8", "#7c5cff", "#0c1020", "#e8ecf6"),
    "lava": ("#ff7a45", "#b3261e", "#15100d", "#f3e9e2"),
    "sandstone": ("#c06b4f", "#6b7a3a", "#fbf4ec", "#3a2f28"),
    "obsidian": ("#c9a24b", "#d9b85a", "#0e0e10", "#ece3d0"),
    "concrete": ("#2d6cff", "#ffd23f", "#f2f2f0", "#161616"),
    "glacier": ("#3a7ca5", "#9ec5d8", "#ffffff", "#22323d"),
}


def rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_bg(slide, hex_bg: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_bg)


def add_title(slide, text, primary, w=13.333):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(w - 1.2), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = rgb(primary)
    return tb


def add_body(slide, text, fg, size=20):
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = rgb(fg)
    return tb


def render(slide, prs, primary, accent, bg, fg):
    layout = prs.slide_layouts[6]  # blank
    s = prs.slides.add_slide(layout)
    set_bg(s, bg)
    t = slide.get("type", "bullets")
    title = slide.get("title", "")
    if t == "title":
        add_title(s, title, primary)
        add_body(s, slide.get("subtitle", ""), fg, 22)
    elif t == "section":
        add_title(s, title, primary)
        add_body(s, slide.get("lead", ""), fg, 20)
    elif t == "bullets":
        add_title(s, title, primary)
        tb = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.2))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for pt in slide.get("points", []):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = "• " + pt
            p.font.size = Pt(22)
            p.font.color.rgb = rgb(fg)
        if slide.get("notes"):
            n = s.notes_slide.notes_text_frame
            n.text = "备注：" + slide["notes"]
    elif t == "two-col":
        add_title(s, title, primary)
        l = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(5.8), Inches(5.0))
        r = s.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.8), Inches(5.0))
        for box, txt in ((l, slide.get("left", "")), (r, slide.get("right", ""))):
            box.text_frame.word_wrap = True
            box.text_frame.paragraphs[0].text = txt
            box.text_frame.paragraphs[0].font.size = Pt(20)
            box.text_frame.paragraphs[0].font.color.rgb = rgb(fg)
    elif t == "quote":
        add_title(s, title, primary)
        add_body(s, "“" + slide.get("text", "") + "”\n— " + slide.get("cite", ""), fg, 24)
    elif t == "steps":
        add_title(s, title, primary)
        tb = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.2))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for i, st in enumerate(slide.get("steps", []), 1):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = f"{i}. {st.get('t','')}：{st.get('d','')}"
            p.font.size = Pt(20)
            p.font.color.rgb = rgb(fg)
    elif t == "table":
        add_title(s, title, primary)
        headers = slide.get("headers", [])
        rows = slide.get("rows", [])
        nr = len(rows) + 1
        nc = max(len(headers), *(len(r) for r in rows), 1)
        gtable = s.shapes.add_table(nr, nc, Inches(0.7), Inches(1.8), Inches(11.9), Inches(0.6 * nr)).table
        for c, h in enumerate(headers):
            cell = gtable.cell(0, c)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(primary)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(18)
            para.font.bold = True
            para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for ri, row in enumerate(rows, 1):
            for ci, val in enumerate(row):
                cell = gtable.cell(ri, ci)
                cell.text = str(val)
                cell.text_frame.paragraphs[0].font.size = Pt(16)
                if ri % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(bg)
    elif t == "chart":
        add_title(s, title, primary)
        chart = slide.get("chart", "bar")
        data = slide.get("data", {})
        labels = data.get("labels", [])
        try:
            values = [float(v) for v in data.get("values", [])]
        except (ValueError, TypeError):
            # 图表数值非法（含非数值字符串）：用 0 占位并保留原长度，
            # 避免单页坏数据让整个 .pptx 构建抛异常崩溃。
            values = [0.0 for _ in data.get("values", [])]
        cd = CategoryChartData()
        cd.categories = labels
        cd.add_series("数值", values)
        ctype = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
                 "pie": XL_CHART_TYPE.PIE,
                 "line": XL_CHART_TYPE.LINE}.get(chart, XL_CHART_TYPE.COLUMN_CLUSTERED)
        gframe = s.shapes.add_chart(ctype, Inches(1.2), Inches(1.9), Inches(11.0), Inches(5.0), cd)
        gframe.chart.has_legend = chart == "pie"
    elif t == "image":
        add_title(s, title, primary)
        if slide.get("src"):
            try:
                s.shapes.add_picture(slide["src"], Inches(3.5), Inches(1.9), Inches(6.3))
            except Exception:
                add_body(s, "[图片加载失败: " + slide["src"] + "]", fg, 16)
        if slide.get("caption"):
            add_body(s, slide["caption"], fg, 16)
    elif t == "closing":
        add_title(s, title, primary)
    else:
        add_title(s, title, primary)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Generate a .pptx deck from a deck spec.")
    ap.add_argument("--spec", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--preset", default=None)
    args = ap.parse_args(argv)

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    pid = args.preset or spec.get("preset") or "academic-light"
    primary, accent, bg, fg = PRESET_COLORS.get(pid, PRESET_COLORS["academic-light"])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for sl in spec.get("slides", []):
        render(sl, prs, primary, accent, bg, fg)
    prs.save(args.out)
    print(f"OK wrote {args.out} ({len(spec.get('slides', []))} slides, preset={pid})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
