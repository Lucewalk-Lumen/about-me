#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""md2docx.py — 将教材 Markdown 转为 docx（含图片嵌入与路径映射）。

前置条件：python-docx 已安装（pip3 install python-docx）。

用法：
  # 单个文件
  python3 md2docx.py <input.md> <output.docx> [--title "文档标题"] \
      [--image-dir <目录> ...]

  # 批量（成对传入）
  python3 md2docx.py --pairs a.md a.docx b.md b.docx [--image-dir <目录> ...]

图片处理关键点：
  markdown 中的图片引用可能是原作者本地路径（当前机器不存在），脚本会按文件名
  （basename）在 --image-dir 指定的目录中递归搜索真实文件并嵌入；找不到时输出
  红色占位提示。转换完成后会打印 嵌入数/引用数，应保证两者一致。
"""

import argparse
import re
import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def find_image(src_path, image_dirs):
    """按文件名在候选目录中查找真实图片路径。"""
    filename = os.path.basename(src_path)
    for search_dir in image_dirs:
        for root, _, files in os.walk(search_dir):
            if filename in files:
                return os.path.join(root, filename)
    return None


def set_cell_shading(cell, color):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    shading_elm.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_formatted_text(paragraph, text):
    """解析行内 markdown 格式（加粗/斜体/行内代码）并添加 runs。"""
    parts = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('***') and part.endswith('***'):
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            paragraph.add_run(part)


def add_image(doc, alt_text, src_path, image_dirs):
    """嵌入图片，找不到真实文件时输出占位提示。"""
    actual_path = find_image(src_path, image_dirs)
    if not actual_path:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'[图片缺失: {alt_text}]')
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        run.italic = True
        return False
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(actual_path, width=Inches(5.5))
        if alt_text and not alt_text.startswith('http'):
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(alt_text)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            run.italic = True
        return True
    except Exception as e:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'[图片加载失败: {alt_text} - {str(e)}]')
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        run.italic = True
        return False


def md_to_docx(md_path, docx_path, image_dirs, title=""):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)

    for i in range(1, 5):
        heading_style = doc.styles[f'Heading {i}']
        heading_style.font.name = 'Microsoft YaHei'
        heading_style.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.size = Pt(22)
        run.bold = True
        doc.add_page_break()

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_buffer = []
    in_table = False
    table_rows = []
    img_count = 0
    img_found = 0

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith('```'):
            if in_code_block:
                code_text = '\n'.join(code_buffer)
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'clear')
                shading.set(qn('w:color'), 'auto')
                shading.set(qn('w:fill'), 'F5F5F5')
                p.paragraph_format.element.get_or_add_pPr().append(shading)
                code_buffer = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', line.strip())
        if img_match:
            alt_text = img_match.group(1)
            src_path = img_match.group(2)
            img_count += 1
            if add_image(doc, alt_text, src_path, image_dirs):
                img_found += 1
            i += 1
            continue

        if '|' in line and line.strip().startswith('|'):
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if all(re.match(r'^[-:]+$', c) for c in cells):
                i += 1
                continue
            if not in_table:
                in_table = True
                table_rows = [cells]
            else:
                table_rows.append(cells)
            if i + 1 >= len(lines) or '|' not in lines[i + 1] or not lines[i + 1].strip().startswith('|'):
                if table_rows:
                    num_cols = max(len(r) for r in table_rows)
                    table = doc.add_table(rows=len(table_rows), cols=num_cols)
                    table.style = 'Table Grid'
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for ri, row_data in enumerate(table_rows):
                        for ci, cell_text in enumerate(row_data):
                            if ci < num_cols:
                                cell = table.cell(ri, ci)
                                cell.text = ''
                                p = cell.paragraphs[0]
                                add_formatted_text(p, cell_text)
                                p.style = doc.styles['Normal']
                                if ri == 0:
                                    set_cell_shading(cell, 'E8E8E8')
                                    for run in p.runs:
                                        run.bold = True
                    doc.add_paragraph()
                in_table = False
                table_rows = []
            i += 1
            continue

        if re.match(r'^---+\s*$', line.strip()) or re.match(r'^\*\*\*+\s*$', line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '999999')
            pBdr.append(bottom)
            p.paragraph_format.element.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        heading_match = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        if line.strip().startswith('>'):
            text = line.strip().lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            add_formatted_text(p, text)
            pBdr = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'), 'single')
            left.set(qn('w:sz'), '12')
            left.set(qn('w:space'), '4')
            left.set(qn('w:color'), '4472C4')
            pBdr.append(left)
            p.paragraph_format.element.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        list_match = re.match(r'^(\s*)[-*]\s+(.+)$', line)
        if list_match:
            indent_level = len(list_match.group(1)) // 2
            text = list_match.group(2)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(1.27 * (indent_level + 1))
            add_formatted_text(p, text)
            i += 1
            continue

        olist_match = re.match(r'^(\s*)\d+\.\s+(.+)$', line)
        if olist_match:
            indent_level = len(olist_match.group(1)) // 2
            text = olist_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Cm(1.27 * (indent_level + 1))
            add_formatted_text(p, text)
            i += 1
            continue

        if line.strip().startswith('<!--'):
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        text = line.strip()
        if text:
            p = doc.add_paragraph()
            add_formatted_text(p, text)

        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")
    print(f"  Images: {img_found}/{img_count} embedded")


def main():
    parser = argparse.ArgumentParser(description='Markdown → docx（含图片嵌入）')
    parser.add_argument('paths', nargs='+', help='md/docx 路径（单个：<in.md> <out.docx>；批量：--pairs 成对）')
    parser.add_argument('--pairs', action='store_true', help='paths 按 md,docx 成对解析')
    parser.add_argument('--title', default='', help='文档标题（可选）')
    parser.add_argument('--image-dir', action='append', default=[], help='图片资源目录（可重复），按文件名搜索嵌入')
    args = parser.parse_args()

    if args.pairs:
        if len(args.paths) % 2 != 0:
            parser.error('--pairs 模式下路径必须成对（md, docx, md, docx ...）')
        pairs = [(args.paths[i], args.paths[i + 1]) for i in range(0, len(args.paths), 2)]
    elif len(args.paths) == 2:
        pairs = [tuple(args.paths)]
    else:
        parser.error('需传入 <input.md> <output.docx> 两个路径，或用 --pairs 批量传入')

    for md_path, docx_path in pairs:
        print(f"\nConverting: {md_path}")
        md_to_docx(md_path, docx_path, args.image_dir, title=args.title)

    print("\nAll done!")


if __name__ == '__main__':
    main()
