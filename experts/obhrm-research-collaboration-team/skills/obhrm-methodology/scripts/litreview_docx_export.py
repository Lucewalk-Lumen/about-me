#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> Word (.docx) converter for OBHRM literature reviews.

将学术 Markdown（尤其中文文献综述风格）转换为排版规范的 Word 文件。
字体、行距、首行缩进、最大标题层级均从 ``../config.json`` 的 ``docx_style`` 读取，
路径与解释器不硬编码。

支持的 Markdown 元素
--------------------
* 标题 ``#``~``####``（上限由 config 决定，更深层级封顶）
* 行内强调 ``**粗体**``、``*斜体*``、`` `代码` ``
* 无序列表（``-``/``*``/``+``）、有序列表（``1.``）
* 分隔线（``---``/``***``/``___``）
* 引用块（``>``）
* 段落（连续非空行合并）
* 图片行内语法 ``![题注](img.png)``：居中插图 + 9pt 居中题注（仅正文段内解析）
* Markdown 表格 ``| a | b |``：解析为原生 Word 表格（首行加粗+浅灰底，单元格左对齐）

CLI 额外能力（向后兼容，缺省无副作用）
------------------------------------
* ``--figure <png>``（可重复）：把指定 PNG 作为图插入到「正文之后、参考文献之前」，
  题注优先取同目录/产出目录 ``框架图数据.json`` 的 ``figure_label``+``title``，
  读不到则回退为 ``图1 算法管理双刃剑效应整合框架``。
* ``--register <md>``（可重复）：把指定 Markdown 表格文件作为表格插入，
  位置紧跟最后一个 ``--figure`` 之后（或正文之后、参考文献之前）。

用法
----
    python litreview_docx_export.py <input.md> [--out output.docx] [--config ../config.json]
    python litreview_docx_export.py <input.md> --figure figures/流程图.png \
        --register figures/图表登记表.md

退出码
------
    0  成功
    2  输入/配置错误（文件不存在、config 解析失败）
"""


# 确保套件包可导入（包目录在上级目录）
import os, sys
_skill_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _skill_dir not in sys.path:
    sys.path.append(_skill_dir)
import argparse
import json
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, Inches
except ImportError:
    print("错误：缺少 python-docx 依赖。请运行：pip install python-docx")
    sys.exit(1)

# 允许导入同目录的 format_references（APA7 规范化）
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from obhrm_research_base.core.safe_path import resolve_config_path, resolve_safe_output  # noqa: E402
try:
    import format_references  # noqa: E402
except Exception:  # 规范化模块缺失时降级（不追加参考文献段）
    format_references = None

CODE_FONT = "Courier New"
REFERENCE_HEADERS = ("参考文献", "references", "bibliography", "引用文献")
FIGURE_META_FILENAME = "框架图数据.json"
# 读不到 figure_label 时的默认题注
DEFAULT_FIGURE_CAPTION = "图1 算法管理双刃剑效应整合框架"

# 模块级样式（由 load_style 填充，单进程 CLI 安全）
_STYLE = {
    "body_font": "Times New Roman",
    "cjk_body_font": "SimSun",
    "heading_font": "Times New Roman",
    "cjk_heading_font": "SimHei",
    "line_spacing": 1.5,
    "first_line_indent_pt": 24,
    "font_size_pt": 12,
    "max_heading_level": 4,
    # 论文格式参数（页边距 cm / 标题字号 pt），缺失时由 load_style 给默认
    "h1_size_pt": 16.0,
    "h2_size_pt": 14.0,
    "h3_size_pt": 12.0,
    "margin_top_cm": 2.54,
    "margin_bottom_cm": 2.54,
    "margin_left_cm": 3.18,
    "margin_right_cm": 3.18,
}

# 行内图片解析时的基准目录（由 markdown_to_docx 写入，缺省为 cwd）
_BASE_DIR = ""


# --------------------------------------------------------------------------- #
# 配置加载
# --------------------------------------------------------------------------- #
def load_style(config_path=None):
    """加载 ../config.json 的 docx_style。

    Args:
        config_path: 显式 ../config.json 路径；为 None 时回退到脚本上级目录的 ../config.json。

    Returns:
        dict: 含 body_font / cjk_body_font / heading_font / cjk_heading_font /
            line_spacing / first_line_indent_pt / font_size_pt / max_heading_level
            的样式字典。

    Raises:
        FileNotFoundError: ../config.json 不存在。
        ValueError: ../config.json 不是合法 JSON。
    """
    if config_path is None:
        config_path = resolve_config_path(config_path, __file__)
    config_path = os.path.abspath(config_path)
    if not os.path.isfile(config_path):
        raise FileNotFoundError("../config.json 未找到: %s" % config_path)
    with open(config_path, encoding="utf-8") as fh:
        cfg = json.load(fh)
    ds = cfg.get("docx_style", {}) or {}
    _STYLE.update({
        "body_font": ds.get("body_font", "Times New Roman"),
        "cjk_body_font": ds.get("cjk_body_font", "SimSun"),
        "heading_font": ds.get("heading_font", "Times New Roman"),
        "cjk_heading_font": ds.get("cjk_heading_font", "SimHei"),
        "line_spacing": float(ds.get("line_spacing", 1.5)),
        "first_line_indent_pt": float(ds.get("first_line_indent_pt", 24)),
        "font_size_pt": float(ds.get("font_size_pt", 12)),
        "max_heading_level": int(ds.get("max_heading_level", 4)),
        "h1_size_pt": float(ds.get("h1_size_pt", 16)),
        "h2_size_pt": float(ds.get("h2_size_pt", 14)),
        "h3_size_pt": float(ds.get("h3_size_pt", 12)),
        "margin_top_cm": float(ds.get("margin_top_cm", 2.54)),
        "margin_bottom_cm": float(ds.get("margin_bottom_cm", 2.54)),
        "margin_left_cm": float(ds.get("margin_left_cm", 3.18)),
        "margin_right_cm": float(ds.get("margin_right_cm", 3.18)),
    })
    return dict(_STYLE)


# --------------------------------------------------------------------------- #
# 字体辅助
# --------------------------------------------------------------------------- #
def _set_font_on_element(element, ascii_font, cjk_font):
    """为 rPr 元素设置 ascii/hAnsi/eastAsia/cs 字体。"""
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), cjk_font)
    rfonts.set(qn("w:cs"), ascii_font)


def set_style_font(style, ascii_font, cjk_font):
    """为 python-docx 样式设置拉丁 + 中日韩字体。"""
    style.font.name = ascii_font
    _set_font_on_element(style.element, ascii_font, cjk_font)


def _set_run_font(run, ascii_font, cjk_font):
    """为单个 run 设置拉丁 + 中日韩字体（含字号以外的字体族）。"""
    run.font.name = ascii_font
    _set_font_on_element(run._element, ascii_font, cjk_font)


def apply_style(document):
    """依据 _STYLE 配置文档基础样式（正文两端对齐 + 标题字体/字号）。"""
    normal = document.styles["Normal"]
    normal.font.size = Pt(_STYLE["font_size_pt"])
    normal.paragraph_format.line_spacing = _STYLE["line_spacing"]
    # 正文两端对齐（中文学术论文常见格式，Justified 使段落左右边缘齐整）
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # 正文：西文 Times New Roman + 中文 宋体(SimSun)
    set_style_font(normal, _STYLE["body_font"], _STYLE["cjk_body_font"])

    # 标题字号：H1≈16pt / H2≈14pt / H3≈12pt，超出 H3 用正文基础字号
    heading_sizes = {
        1: _STYLE["h1_size_pt"],
        2: _STYLE["h2_size_pt"],
        3: _STYLE["h3_size_pt"],
    }
    for level in range(1, _STYLE["max_heading_level"] + 1):
        style_name = "Heading %d" % level
        if style_name in document.styles:
            style = document.styles[style_name]
            size = heading_sizes.get(level, _STYLE["font_size_pt"])
            style.font.size = Pt(size)
            style.font.bold = True
            # 标题：西文 Times New Roman + 中文 黑体(SimHei) + 加粗
            set_style_font(style, _STYLE["heading_font"], _STYLE["cjk_heading_font"])


def _apply_page_margins(document):
    """设置论文页边距：上下 2.54cm、左右 3.18cm（从 config 读取，缺失用默认）。"""
    section = document.sections[0]
    section.top_margin = Cm(_STYLE["margin_top_cm"])
    section.bottom_margin = Cm(_STYLE["margin_bottom_cm"])
    section.left_margin = Cm(_STYLE["margin_left_cm"])
    section.right_margin = Cm(_STYLE["margin_right_cm"])


# --------------------------------------------------------------------------- #
# 行内解析
# --------------------------------------------------------------------------- #
_INLINE_RE = re.compile(
    r"`([^`]+?)`"          # group 1: 行内代码
    r"|\*\*(.+?)\*\*"       # group 2: 粗体
    r"|\*(.+?)\*"           # group 3: 斜体
)
# 行内图片：![alt](path)
_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def parse_inline(text):
    """将一行文本拆分为带样式的 run 列表。"""
    runs = []
    last = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > last:
            runs.append((text[last:match.start()], {}))
        if match.group(1) is not None:
            runs.append((match.group(1), {"code": True}))
        elif match.group(2) is not None:
            runs.append((match.group(2), {"bold": True}))
        elif match.group(3) is not None:
            runs.append((match.group(3), {"italic": True}))
        last = match.end()
    if last < len(text):
        runs.append((text[last:], {}))
    return runs


def add_styled_runs(paragraph, text):
    """将解析后的行内内容作为 run 添加到段落。"""
    for run_text, style in parse_inline(text):
        if run_text == "":
            continue
        run = paragraph.add_run(run_text)
        if style.get("bold"):
            run.bold = True
        if style.get("italic"):
            run.italic = True
        if style.get("code"):
            run.font.name = CODE_FONT
            _set_font_on_element(run._element, CODE_FONT, CODE_FONT)


# --------------------------------------------------------------------------- #
# 块级构建（标题 / 段落 / 列表 / 引用 / 分隔线）
# --------------------------------------------------------------------------- #
def add_heading(document, text, level):
    """添加标题段落（层级封顶到 max_heading_level）。"""
    level = max(1, min(level, _STYLE["max_heading_level"]))
    paragraph = document.add_paragraph(style="Heading %d" % level)
    add_styled_runs(paragraph, text)
    return paragraph


def add_paragraph(document, text, first_line_indent=True):
    """添加正文段落（默认首行缩进，两端对齐）。"""
    paragraph = document.add_paragraph()
    # 正文两端对齐（中文学术论文常见格式，Justified 使段落左右边缘齐整）
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Pt(_STYLE["first_line_indent_pt"])
    add_styled_runs(paragraph, text)


def add_bullet_list(document, items):
    """添加无序列表。"""
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        add_styled_runs(paragraph, item)


def add_numbered_list(document, items):
    """添加有序列表。"""
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        add_styled_runs(paragraph, item)


def add_blockquote(document, text):
    """添加引用块（缩进 + 斜体，保持左对齐）。"""
    paragraph = document.add_paragraph()
    # 引用块保持左对齐，避免继承 Normal 的两端对齐导致引用文字拉伸难看
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Pt(20)
    paragraph.paragraph_format.right_indent = Pt(20)
    run = paragraph.add_run(text)
    run.italic = True


def add_horizontal_rule(document):
    """添加仅含下边框的分隔段落。"""
    paragraph = document.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


# --------------------------------------------------------------------------- #
# 图片 / 表格（Word 原生嵌入）
# --------------------------------------------------------------------------- #
def _resolve_image(path, base_dirs):
    """解析图片/文件相对路径，返回绝对路径或 None。

    解析顺序：绝对路径直接判定 -> 依次拼接各基准目录 -> 当前工作目录相对路径。
    """
    if os.path.isabs(path) and os.path.isfile(path):
        return os.path.abspath(path)
    candidates = []
    for d in base_dirs:
        if d:
            candidates.append(os.path.join(d, path))
    candidates.append(path)  # 相对 cwd 的兜底
    for cand in candidates:
        if os.path.isfile(cand):
            return os.path.abspath(cand)
    return None


def _figure_width(document):
    """计算图片宽度：页面内容宽度的 75% 与 6.0in 的较小者（单位 Emu）。"""
    section = document.sections[0]
    page_w = section.page_width or Cm(21.59)
    left = section.left_margin or Cm(_STYLE["margin_left_cm"])
    right = section.right_margin or Cm(_STYLE["margin_right_cm"])
    content_emu = int(page_w) - int(left) - int(right)
    if content_emu < 1:
        content_emu = int(Cm(21.59 - _STYLE["margin_left_cm"]
                             - _STYLE["margin_right_cm"]))
    candidate = int(content_emu * 0.75)
    cap = int(Inches(6.0))
    return min(candidate, cap)


def _add_figure(document, image_path, caption_text):
    """居中插入图片并加 9pt 居中题注，返回 (图片段落元素, 题注段落元素)。

    Args:
        document: python-docx Document。
        image_path: 已解析的绝对图片路径。
        caption_text: 题注文本（如「图1 …」）。

    Returns:
        tuple: (图片所在 <w:p> 元素, 题注所在 <w:p> 元素)。
    """
    pic_para = document.add_paragraph()
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_para.add_run()
    run.add_picture(image_path, width=_figure_width(document))

    cap_para = document.add_paragraph()
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_para.paragraph_format.space_before = Pt(4)
    cap_run = cap_para.add_run(caption_text)
    cap_run.font.size = Pt(9)
    # 题注继承正文字体族（西文 + 中文）
    _set_run_font(cap_run, _STYLE["body_font"], _STYLE["cjk_body_font"])
    return pic_para._p, cap_para._p


def _emit_image_line(document, line):
    """处理一个可能含行内图片语法的行：按文本/图片分段落依次输出。

    图片前/后的纯文本作为普通段落；图片本身居中插入并加 9pt 题注。
    """
    base_dirs = [_BASE_DIR, os.getcwd()]
    last = 0
    for match in _IMAGE_RE.finditer(line):
        pre = line[last:match.start()].strip()
        if pre:
            add_paragraph(document, pre)
        alt = match.group(1).strip()
        img = match.group(2).strip()
        resolved = _resolve_image(img, base_dirs)
        if resolved:
            _add_figure(document, resolved, alt if alt else os.path.basename(resolved))
        else:
            print("警告: 图片未找到，跳过: %s" % img, file=sys.stderr)
        last = match.end()
    post = line[last:].strip()
    if post:
        add_paragraph(document, post)


def _set_cell_shading(cell, fill_hex):
    """为表格单元格设置底纹颜色（如浅灰 'D9D9D9'）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _is_separator_row(line):
    """判断一行是否为 Markdown 表分隔行（如 |---|:---:|）。

    去掉所有竖线后再判定，避免表内竖线干扰。
    """
    s = line.strip()
    if "|" not in s:
        return False
    body = s.replace("|", "").strip()
    if not body:
        return False
    if "-" not in body:
        return False
    return all(ch in "-: " for ch in body)


def _is_table_row(line):
    """判断一行是否像 Markdown 表格数据/表头行。"""
    s = line.strip()
    if "|" not in s:
        return False
    return s.startswith("|") or s.endswith("|")


def _split_row(line):
    """将表格行拆分为单元格文本列表（去除首尾竖线并去空白）。"""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [cell.strip() for cell in s.split("|")]


def _build_table(document, rows):
    """依据二维文本列表构建原生 Word 表格。

    首行（表头）加粗 + 浅灰底；所有单元格左对齐；使用 Table Grid 默认边框。

    Args:
        document: python-docx Document。
        rows: 列表的列表，rows[0] 为表头，后续为数据行（已去除分隔行）。

    Returns:
        w:tbl: 新建表格的底层 XML 元素，便于后续位置调整。
    """
    if not rows:
        return None
    n_cols = max(len(r) for r in rows)
    table = document.add_table(rows=0, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx in range(n_cols):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = cells[c_idx]
            cell.text = ""  # 清空默认空 run
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(text)
            if r_idx == 0:
                run.bold = True
                _set_cell_shading(cell, "D9D9D9")
    return table._tbl


def _extract_all_tables(md_text):
    """从 Markdown 文本中提取所有表格，返回表格行列表的列表。

    每个表格表示为 ``rows``（含表头，不含分隔行）。忽略标题、引用块等非表格内容。
    """
    lines = md_text.splitlines()
    total = len(lines)
    tables = []
    idx = 0
    while idx < total:
        line = lines[idx].strip()
        # 表头行 + 下一行为分隔行 => 命中一个表格
        if _is_table_row(line) and idx + 1 < total and _is_separator_row(lines[idx + 1].strip()):
            rows = [_split_row(line)]
            idx += 2  # 跳过表头与分隔行
            while idx < total and _is_table_row(lines[idx].strip()):
                rows.append(_split_row(lines[idx].strip()))
                idx += 1
            tables.append(rows)
        else:
            idx += 1
    return tables


def _is_table_start(lines, index):
    """判断 lines[index] 是否为一个表格的起始（表头行且下一行为分隔行）。"""
    if not _is_table_row(lines[index].strip()):
        return False
    if index + 1 < len(lines) and _is_separator_row(lines[index + 1].strip()):
        return True
    return False


def _parse_table_at(lines, index):
    """从 index 处解析一个表格，返回 (rows, 新 index)。"""
    rows = [_split_row(lines[index].strip())]
    index += 2  # 表头 + 分隔行
    while index < len(lines) and _is_table_row(lines[index].strip()):
        rows.append(_split_row(lines[index].strip()))
        index += 1
    return rows, index


# --------------------------------------------------------------------------- #
# 命令行 --figure / --register 相关
# --------------------------------------------------------------------------- #
def _load_figure_meta(search_dirs):
    """在 search_dirs 中查找 框架图数据.json，返回解析后的 dict 或 None。"""
    for d in search_dirs:
        if not d:
            continue
        cand = os.path.join(d, FIGURE_META_FILENAME)
        if os.path.isfile(cand):
            try:
                with open(cand, encoding="utf-8") as fh:
                    return json.load(fh)
            except (ValueError, OSError):
                return None
    return None


def _figure_caption(meta, fallback_name):
    """根据 框架图数据.json 构造图的题注文本。

    优先用 ``figure_label`` + ``title``；读不到则回退为默认题注。
    """
    if isinstance(meta, dict):
        label = (meta.get("figure_label") or "").strip()
        title = (meta.get("title") or "").strip()
        parts = [p for p in (label, title) if p]
        if parts:
            return " ".join(parts)
    return DEFAULT_FIGURE_CAPTION


def _find_reference_anchor(document):
    """查找参考文献区段标题段落（Heading 且文本命中 REFERENCE_HEADERS）。"""
    for para in document.paragraphs:
        if (para.style is not None
                and para.style.name.startswith("Heading")
                and _is_reference_header(para.text)):
            return para
    for para in document.paragraphs:
        if _is_reference_header(para.text):
            return para
    return None


def _reposition_blocks_before(document, elements, anchor):
    """将若干块级元素（<w:p>/<w:tbl>）移动到 anchor 段落之前；anchor 为 None 则追加到末尾。"""
    body = document.element.body
    for el in elements:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
    if anchor is None:
        for el in elements:
            body.append(el)
        return
    anchor_el = anchor._p
    children = list(body)
    if anchor_el not in children:
        for el in elements:
            body.append(el)
        return
    pos = children.index(anchor_el)
    for i, el in enumerate(elements):
        body.insert(pos + i, el)


def _insert_figures_register(document, figure_paths, register_paths, md_dir):
    """插入 CLI 指定的图片与登记表表格，位置统一在「正文之后、参考文献之前」。

    Args:
        document: python-docx Document。
        figure_paths: ``--figure`` 列表（可重复），可为 None。
        register_paths: ``--register`` 列表（可重复），可为 None。
        md_dir: 输入 Markdown 所在目录，用于相对路径与 框架图数据.json 查找。
    """
    new_blocks = []
    base_dirs = [os.getcwd(), md_dir]

    if figure_paths:
        for fig in figure_paths:
            resolved = _resolve_image(fig, base_dirs)
            if not resolved:
                print("警告: 图片未找到，跳过: %s" % fig, file=sys.stderr)
                continue
            meta = _load_figure_meta(
                [os.path.dirname(resolved), md_dir, os.getcwd()])
            caption = _figure_caption(meta, os.path.basename(resolved))
            pic_p, cap_p = _add_figure(document, resolved, caption)
            new_blocks.extend([pic_p, cap_p])

    if register_paths:
        for reg in register_paths:
            resolved = _resolve_image(reg, base_dirs)
            if not resolved:
                print("警告: 登记表未找到，跳过: %s" % reg, file=sys.stderr)
                continue
            with open(resolved, encoding="utf-8") as fh:
                text = fh.read()
            tables = _extract_all_tables(text)
            if not tables:
                print("警告: 登记表未解析到表格，跳过: %s" % reg, file=sys.stderr)
                continue
            # 给登记表加一个小标题，便于正文引用
            heading = add_heading(document, "图表登记表", 2)
            new_blocks.append(heading._p)
            for rows in tables:
                if not rows:
                    continue
                tbl = _build_table(document, rows)
                if tbl is not None:
                    new_blocks.append(tbl)

    if not new_blocks:
        return
    anchor = _find_reference_anchor(document)
    _reposition_blocks_before(document, new_blocks, anchor)


# --------------------------------------------------------------------------- #
# 参考文献
# --------------------------------------------------------------------------- #
def load_reference_paragraphs(path):
    """从 APA7 参考文献 Markdown 读取条目（跳过标题行，收集非空段落）。"""
    paras = []
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            s = line.strip()
            if not s or s.startswith("#"):
                continue
            if s.startswith("- "):
                s = s[2:].strip()
            paras.append(s)
    return paras


def add_references(document, paragraphs):
    """追加 APA7 参考文献段（悬挂缩进：首行齐左、续行缩进）。"""
    add_heading(document, "参考文献", 1)
    for para in paragraphs:
        p = document.add_paragraph()
        # 参考文献保持左对齐（APA7 悬挂缩进通常为左对齐，不被两端对齐破坏排版）
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-24)
        add_styled_runs(p, para)


# --------------------------------------------------------------------------- #
# Markdown 解析调度
# --------------------------------------------------------------------------- #
HR_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
UNORDERED_RE = re.compile(r"^[-*+]\s+(.*)$")
ORDERED_RE = re.compile(r"^\d+\.\s+(.*)$")
QUOTE_RE = re.compile(r"^>\s?(.*)$")


def _is_reference_header(text):
    """判断文本是否像参考文献区段标题。"""
    lowered = text.strip().lower()
    return any(token in lowered for token in REFERENCE_HEADERS)


def convert(md_text):
    """解析 Markdown 文本，返回已排版的 Document。

    Args:
        md_text: Markdown 源文本。

    Returns:
        docx.Document: 已填充样式的文档对象。
    """
    document = Document()
    apply_style(document)
    _apply_page_margins(document)
    lines = md_text.splitlines()

    in_references = False
    index = 0
    total = len(lines)

    while index < total:
        raw_line = lines[index]
        stripped = raw_line.strip()
        para_parts = []  # 段落部件初始化

        if stripped == "":
            index += 1
            continue

        if HR_RE.match(stripped):
            add_horizontal_rule(document)
            index += 1
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match is not None:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            add_heading(document, text, level)
            if _is_reference_header(text):
                in_references = True
            index += 1
            continue

        if stripped.startswith(">"):
            quote_parts = []
            while index < total and QUOTE_RE.match(lines[index].strip()):
                quote_text = QUOTE_RE.match(lines[index].strip()).group(1).strip()
                quote_parts.append(quote_text)
                index += 1
            add_blockquote(document, " ".join(quote_parts))
            continue

        if not in_references:
            # 行内图片：![alt](path) -> 居中插图 + 9pt 题注
            if _IMAGE_RE.search(stripped):
                if para_parts:
                    add_paragraph(document, " ".join(para_parts))
                    para_parts = []
                _emit_image_line(document, stripped)
                index += 1
                continue

            # Markdown 表格 -> 原生 Word 表格
            if _is_table_start(lines, index):
                if para_parts:
                    add_paragraph(document, " ".join(para_parts))
                    para_parts = []
                rows, index = _parse_table_at(lines, index)
                _build_table(document, rows)
                continue

            if UNORDERED_RE.match(stripped):
                items = []
                while index < total:
                    current = lines[index].strip()
                    match = UNORDERED_RE.match(current)
                    if match is not None:
                        items.append(match.group(1))
                        index += 1
                    elif current == "":
                        break
                    else:
                        break
                add_bullet_list(document, items)
                continue

            if ORDERED_RE.match(stripped):
                items = []
                while index < total:
                    current = lines[index].strip()
                    match = ORDERED_RE.match(current)
                    if match is not None:
                        items.append(match.group(1))
                        index += 1
                    elif current == "":
                        break
                    else:
                        break
                add_numbered_list(document, items)
                continue

        # 普通段落累积
        para_parts = []
        while index < total:
            current = lines[index].strip()
            if current == "":
                break
            if HR_RE.match(current):
                break
            if HEADING_RE.match(current):
                break
            if current.startswith(">"):
                break
            if not in_references and (
                UNORDERED_RE.match(current) or ORDERED_RE.match(current)
                or _IMAGE_RE.search(current) or _is_table_start(lines, index)
            ):
                break
            para_parts.append(current)
            index += 1
        add_paragraph(document, " ".join(para_parts))

    return document


# --------------------------------------------------------------------------- #
# 入口
# --------------------------------------------------------------------------- #
def markdown_to_docx(input_path, output_path=None, config_path=None,
                     references_paragraphs=None, figure_paths=None,
                     register_paths=None):
    """将 Markdown 文件转换为 Word 文档。

    Args:
        input_path: 源 .md 路径。
        output_path: 可选输出 .docx 路径；None 时与输入同目录同名 .docx。
        config_path: 可选 ../config.json 路径。
        references_paragraphs: 可选 APA7 参考文献条目列表，非空时追加参考文献段。
        figure_paths: 可选 ``--figure`` PNG 路径列表，插入正文后、参考文献前。
        register_paths: 可选 ``--register`` Markdown 表格路径列表，紧跟 figure 之后。

    Returns:
        str: 生成的 .docx 绝对路径。

    Raises:
        FileNotFoundError: 输入或 config 不存在。
        ValueError: ../config.json 解析失败。
    """
    load_style(config_path)  # 读取样式
    input_path = os.path.abspath(input_path)
    global _BASE_DIR
    _BASE_DIR = os.path.dirname(input_path)
    if not os.path.isfile(input_path):
        raise FileNotFoundError("Markdown 输入未找到: %s" % input_path)

    if output_path is None:
        base, _ = os.path.splitext(input_path)
        output_path = base + ".docx"
    # 强制 范围校验，防止输出路径写出允许范围
    output_path = str(resolve_safe_output(output_path, base=os.getcwd()))

    with open(input_path, "r", encoding="utf-8") as handle:
        md_text = handle.read()

    document = convert(md_text)
    if references_paragraphs:
        add_references(document, references_paragraphs)
    # 图片 / 登记表：统一插入到正文之后、参考文献之前（或无参考文献时文档末尾）
    _insert_figures_register(document, figure_paths, register_paths, _BASE_DIR)
    document.save(output_path)
    return output_path


def main(argv):
    """命令行入口。"""
    ap = argparse.ArgumentParser(description="Markdown -> Word 转换器（读 ../config.json 样式）")
    ap.add_argument("input", help="输入 Markdown 文件路径")
    ap.add_argument("--out", default=None, help="输出 .docx 路径（默认与输入同名）")
    ap.add_argument("--config", default=None, help="../config.json 路径（默认脚本上级目录）")
    ap.add_argument("--references", default=None,
                    help="已格式化的 APA7 参考文献 Markdown（追加到文末）")
    ap.add_argument("--format-references", action="store_true", default=False,
                    help="自动将 READINGS 归一为 APA7 并追加参考文献段")
    ap.add_argument("--readings", default=None,
                    help="--format-references 的书目源（默认取输入同目录 READINGS.md）")
    ap.add_argument("--figure", action="append", default=None,
                    metavar="PNG",
                    help="将指定 PNG 作为图插入正文后、参考文献前（可重复）")
    ap.add_argument("--register", action="append", default=None,
                    metavar="MD",
                    help="将指定 Markdown 表格文件作为表格插入（紧跟 --figure 之后，可重复）")
    args = ap.parse_args(argv)

    ref_paras = None
    if args.references:
        if not os.path.isfile(args.references):
            print("错误: 参考文献文件未找到: %s" % args.references, file=sys.stderr)
            return 2
        ref_paras = load_reference_paragraphs(args.references)
    elif args.format_references:
        if format_references is None:
            print("错误: format_references 模块不可用", file=sys.stderr)
            return 2
        bib = args.readings
        if not bib:
            cand = os.path.join(os.path.dirname(os.path.abspath(args.input)), "READINGS.md")
            bib = cand if os.path.isfile(cand) else None
        if not bib or not os.path.isfile(bib):
            print("错误: 未找到书目源（用 --readings 指定 READINGS.md）", file=sys.stderr)
            return 2
        with open(bib, encoding="utf-8") as handle:
            ref_paras = format_references.normalize_references(
                handle.read(), format_references._load_reference_style_spec())

    try:
        result_path = markdown_to_docx(
            args.input, args.out, args.config, ref_paras,
            args.figure, args.register)
    except FileNotFoundError as exc:
        print("错误: %s" % exc, file=sys.stderr)
        return 2
    except ValueError as exc:
        print("错误: ../config.json 解析失败: %s" % exc, file=sys.stderr)
        return 2

    size_bytes = os.path.getsize(result_path)
    print("Generated: %s" % result_path)
    print("Size: %.1f KB (%d bytes)" % (size_bytes / 1024.0, size_bytes))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
